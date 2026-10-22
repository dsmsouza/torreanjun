import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\Anjun_\Downloads\levantamento de mds que fecham as 18hrs.xlsx')
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"--- Sheet: {sheetname} ({ws.max_row} rows, {ws.max_column} cols) ---")
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    print("  Header:", headers)
    # Check blank count in each column
    col_blanks = {}
    for c in range(1, ws.max_column + 1):
        h = headers[c-1] or f"Col_{c}"
        blanks = 0
        total = 0
        for r in range(2, ws.max_row + 1):
            val = ws.cell(r, c).value
            total += 1
            if val is None or str(val).strip() == "":
                blanks += 1
        col_blanks[h] = (blanks, total)
    print("  Blanks / Total:", col_blanks)

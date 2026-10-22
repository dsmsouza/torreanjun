import openpyxl
from datetime import datetime
from collections import Counter

wb = openpyxl.load_workbook(r'C:\ANJUN\Receber ordens-20261022080746016-7aPEpU.xlsx', read_only=True)
sheet = wb.active

dates_criacao = Counter()
dates_coleta = Counter()
dates_conclusao = Counter()

# Find col indexes
header = [cell for cell in next(sheet.iter_rows(values_only=True))]
idx_criacao = header.index('TEMPO DE CRIAÇÃO DO PEDIDO') if 'TEMPO DE CRIAÇÃO DO PEDIDO' in header else -1
idx_coleta = header.index('Tempo real para a coleta') if 'Tempo real para a coleta' in header else -1
idx_conclusao = header.index('Data de conclusão') if 'Data de conclusão' in header else -1
idx_status = header.index('Status do pedido') if 'Status do pedido' in header else -1
idx_md = header.index('ponto de coleção') if 'ponto de coleção' in header else -1

print(f"Col indexes: criacao={idx_criacao}, coleta={idx_coleta}, conclusao={idx_conclusao}, status={idx_status}, md={idx_md}")

count = 0
for row in sheet.iter_rows(values_only=True):
    count += 1
    if count == 1:
        continue
    c_coleta = str(row[idx_coleta])[:10] if row[idx_coleta] else "Vazio"
    dates_coleta[c_coleta] += 1
    
    if count % 100000 == 0:
        print(f"Read {count} rows...")

print(f"Total rows: {count}")
print("Top dates in coleta:")
for d, c in dates_coleta.most_common(25):
    print(f"  {d}: {c}")

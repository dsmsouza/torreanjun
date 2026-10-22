import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\Anjun_\Downloads\levantamento de mds que fecham as 18hrs.xlsx', data_only=True)
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print(f"Sheet {sheetname}: EMPTY")
        continue
    header = rows[0]
    # Trim trailing Nones in header
    last_idx = 0
    for idx, h in enumerate(header):
        if h is not None:
            last_idx = idx
    header = header[:last_idx+1]
    
    print(f"\nSheet '{sheetname}': {len(rows)-1} data rows, {len(header)} columns")
    print(f"  Headers: {header}")
    
    # Check blanks per column
    for col_idx, col_name in enumerate(header):
        vals = [r[col_idx] if col_idx < len(r) else None for r in rows[1:]]
        blanks = sum(1 for v in vals if v is None or str(v).strip() == "")
        non_blanks = len(vals) - blanks
        sample = [v for v in vals if v is not None and str(v).strip() != ""][:3]
        print(f"    Col {col_idx+1} '{col_name}': {blanks} blanks / {len(vals)} total. Sample non-blank: {sample}")

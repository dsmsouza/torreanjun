# -*- coding: utf-8 -*-
"""
saturacao_rotas_service.py — Serviço de Cálculo de Saturação de Rotas VUC e Cruzamento com D-1
Cruza o Planejamento de Rotas da Semana com as Bipagens Reais do Dia Anterior (D-1) da Torre ANJUN.
"""

import os
import openpyxl
import pandas as pd
import sqlite3
from typing import Dict, Any, Optional

DB_PATH = r"C:\ANJUN\Banco de Dados\torre_anjun.db"
DEFAULT_EXCEL_PATH = r"C:\Users\Anjun_\Downloads\Planejamento de rotas - MD - 26.10.2026.xlsx"


def obter_coordenadas_mds(conn: sqlite3.Connection) -> Dict[str, Dict[str, Any]]:
    """Carrega mapa completo e deduplicado de coordenadas para todos os MDs."""
    df_pc = pd.read_sql_query(
        "SELECT iata, endere_o, latitude, longitude, cidade FROM pontos_coleta ORDER BY CASE WHEN cidade LIKE '%sao paulo%' OR cidade LIKE '%são paulo%' THEN 0 ELSE 1 END", 
        conn
    )
    df_pc = df_pc.drop_duplicates(subset=['iata'], keep='first')

    df_mc = pd.read_sql_query(
        "SELECT md_codigo as iata, endereco, latitude, longitude, cidade FROM malha_coordenadas ORDER BY CASE WHEN cidade LIKE '%sao paulo%' OR cidade LIKE '%são paulo%' THEN 0 ELSE 1 END", 
        conn
    )
    df_mc = df_mc.drop_duplicates(subset=['iata'], keep='first')

    coords_map = {}
    for _, r in df_pc.iterrows():
        iata = str(r['iata']).strip().upper()
        if pd.notna(r['latitude']) and pd.notna(r['longitude']):
            coords_map[iata] = {'lat': float(r['latitude']), 'lng': float(r['longitude']), 'end': str(r['endere_o'] or '')}

    for _, r in df_mc.iterrows():
        iata = str(r['iata']).strip().upper()
        if pd.notna(r['latitude']) and pd.notna(r['longitude']):
            if iata not in coords_map:
                coords_map[iata] = {'lat': float(r['latitude']), 'lng': float(r['longitude']), 'end': str(r['endereco'] or '')}

    # Correções manuais verificadas para pontos centrais
    coords_map['SP-W-MD074'] = {'lat': -23.5464385, 'lng': -46.6245424, 'end': 'Av. Rangel Pestana, 1013, Brás, São Paulo - SP'}
    coords_map['SP-W-MD207'] = {'lat': -23.531037, 'lng': -46.616824, 'end': 'R. Padre Lima, 184 - Canindé, São Paulo - SP'}
    coords_map['MD-GAS-02'] = {'lat': -23.5556725, 'lng': -46.6143409, 'end': 'Rua do Gasômetro, 539, Brás, São Paulo - SP'}

    return coords_map


def calcular_saturacao_rotas(
    excel_path: Optional[str] = None,
    dia_planejado: str = "26/10/2026",
    dia_anterior: Optional[str] = None,
    capacidade_vuc: int = 2500
) -> Dict[str, Any]:
    excel_file = excel_path or DEFAULT_EXCEL_PATH
    if not os.path.exists(excel_file):
        raise FileNotFoundError(f"Planilha de planejamento não encontrada: {excel_file}")

    conn = sqlite3.connect(f"file:{DB_PATH}?mode=ro", uri=True)

    # 1. Determina data de D-1 (dia anterior com operações registradas)
    if not dia_anterior:
        cur = conn.cursor()
        cur.execute("SELECT DISTINCT dia_base FROM operacional_pedidos_completo ORDER BY rowid DESC LIMIT 1;")
        row = cur.fetchone()
        dia_anterior = str(row[0]) if row else "24/10/2026"

    # 2. Carrega mapa de coordenadas
    coords_map = obter_coordenadas_mds(conn)

    # 3. Carrega volume real bipado no D-1
    q_d1 = """
        SELECT UPPER(TRIM(ponto_colecao)) as md_clean, COUNT(*) as vol_d1
        FROM operacional_pedidos_completo
        WHERE dia_base = ?
        GROUP BY UPPER(TRIM(ponto_colecao));
    """
    df_d1 = pd.read_sql_query(q_d1, conn, params=(dia_anterior,))
    dict_d1 = dict(zip(df_d1['md_clean'], df_d1['vol_d1']))

    # Também checa detalhe_md_historico como backup para D-1
    q_d1_hist = """
        SELECT UPPER(TRIM(md)) as md_clean, SUM(total) as vol_hist
        FROM detalhe_md_historico
        WHERE dia_base = ?
        GROUP BY UPPER(TRIM(md));
    """
    df_d1_hist = pd.read_sql_query(q_d1_hist, conn, params=(dia_anterior,))
    dict_d1_hist = dict(zip(df_d1_hist['md_clean'], df_d1_hist['vol_hist']))

    conn.close()

    # 4. Carrega Planilha de Rotas
    wb = openpyxl.load_workbook(excel_file, read_only=True, data_only=True)

    # Volumetria D-0
    ws_vol = wb['Volumetria']
    vol_rows = list(ws_vol.iter_rows(values_only=True))
    df_vol = pd.DataFrame(vol_rows[1:], columns=vol_rows[0])
    df_vol = df_vol[df_vol['MD'].notna() & df_vol['Rota'].notna()].copy()
    df_vol['md_clean'] = df_vol['MD'].astype(str).str.strip().str.upper()
    df_vol['bipada_d0'] = pd.to_numeric(df_vol['Quantidade bipada'], errors='coerce').fillna(0)
    df_vol['backlog_d0'] = pd.to_numeric(df_vol['Backlog'], errors='coerce').fillna(0)
    df_vol['vol_d0'] = df_vol['bipada_d0'] + df_vol['backlog_d0']

    # Rota da Semana
    ws_rot = wb['Rota da Semana']
    rot_rows = list(ws_rot.iter_rows(values_only=True))
    df_rot = pd.DataFrame(rot_rows[1:], columns=rot_rows[0])
    df_rot = df_rot[df_rot['MD'].notna() & df_rot['Rota'].notna()].copy()
    wb.close()

    # Mapeamento e limpeza de colunas
    col_map = {}
    for c in df_rot.columns:
        if c is None: continue
        c_str = str(c)
        if 'endere' in c_str.lower(): col_map[c] = 'endereco'
        elif 'bairro' in c_str.lower(): col_map[c] = 'bairro'
        elif 'regi' in c_str.lower(): col_map[c] = 'regiao'
        elif 'cidade' in c_str.lower(): col_map[c] = 'cidade'
        elif 'cep' in c_str.lower(): col_map[c] = 'cep'
        elif 'rota' in c_str.lower(): col_map[c] = 'rota'
        elif 'saca' in c_str.lower(): col_map[c] = 'sacas'
        elif 'motorista' in c_str.lower(): col_map[c] = 'motorista'
        elif 'placa' in c_str.lower(): col_map[c] = 'placa'
        elif 'tipo de carro' in c_str.lower() or 'veiculo' in c_str.lower(): col_map[c] = 'tipo_carro'
        elif 'janela 1' in c_str.lower(): col_map[c] = 'janela1'
        elif 'janela 2' in c_str.lower(): col_map[c] = 'janela2'
        elif 'limite' in c_str.lower(): col_map[c] = 'horario_limite'
        elif c_str.strip() == 'MD': col_map[c] = 'md'

    df_rot = df_rot.rename(columns=col_map)
    df_rot['md_clean'] = df_rot['md'].astype(str).str.strip().str.upper()
    df_rot['rota_clean'] = df_rot['rota'].astype(str).str.strip()
    df_rot['sacas'] = pd.to_numeric(df_rot.get('sacas', 0), errors='coerce').fillna(0)
    df_rot['seq'] = df_rot.groupby('rota_clean').cumcount() + 1

    # Merge de Volumetria
    merged = pd.merge(df_rot, df_vol[['md_clean', 'bipada_d0', 'backlog_d0', 'vol_d0']], on='md_clean', how='left')
    merged['vol_d0'] = merged['vol_d0'].fillna(0)
    merged['bipada_d0'] = merged['bipada_d0'].fillna(0)
    merged['backlog_d0'] = merged['backlog_d0'].fillna(0)

    # Vincula volume D-1
    merged['vol_d1'] = merged['md_clean'].apply(lambda x: dict_d1.get(x, dict_d1_hist.get(x, 0)))

    # Vincula Coordenadas
    merged['lat'] = merged['md_clean'].map(lambda x: coords_map.get(x, {}).get('lat'))
    merged['lng'] = merged['md_clean'].map(lambda x: coords_map.get(x, {}).get('lng'))
    merged['endereco_db'] = merged['md_clean'].map(lambda x: coords_map.get(x, {}).get('end', ''))
    merged['endereco_final'] = merged['endereco'].fillna(merged['endereco_db']).fillna('')

    # 5. Agrega por Rota
    rotas_list = []
    for rota_name, group in merged.groupby('rota_clean'):
        tot_vol_d0 = int(group['vol_d0'].sum())
        tot_bip_d0 = int(group['bipada_d0'].sum())
        tot_bk_d0 = int(group['backlog_d0'].sum())
        tot_vol_d1 = int(group['vol_d1'].sum())
        tot_sacas = int(group['sacas'].sum())
        num_mds = len(group)

        motorista = str(group['motorista'].dropna().iloc[0]) if len(group['motorista'].dropna()) > 0 else 'Não atribuído'
        placa = str(group['placa'].dropna().iloc[0]) if len(group['placa'].dropna()) > 0 else 'Sem Placa'
        tipo_carro = str(group['tipo_carro'].dropna().iloc[0]) if 'tipo_carro' in group and len(group['tipo_carro'].dropna()) > 0 else 'VUC'
        regiao = str(group['regiao'].dropna().iloc[0]) if len(group['regiao'].dropna()) > 0 else 'SP'
        cidade = str(group['cidade'].dropna().iloc[0]) if len(group['cidade'].dropna()) > 0 else 'São Paulo'

        # Saturação D-0 (Estimada)
        ocupacao_d0 = round((tot_vol_d0 / capacidade_vuc) * 100, 1) if capacidade_vuc > 0 else 0
        if ocupacao_d0 > 100.0:
            status_d0 = "CRÍTICO"
            badge_d0 = "🔴 CRÍTICO (>100%)"
            cor_hex = "#ef4444"
        elif ocupacao_d0 >= 80.0:
            status_d0 = "ALERTA"
            badge_d0 = "🟠 ALERTA (80-100%)"
            cor_hex = "#f97316"
        elif ocupacao_d0 >= 50.0:
            status_d0 = "EQUILIBRADO"
            badge_d0 = "🟢 EQUILIBRADO (50-80%)"
            cor_hex = "#22c55e"
        elif ocupacao_d0 >= 20.0:
            status_d0 = "MODERADO"
            badge_d0 = "🔵 MODERADO (20-50%)"
            cor_hex = "#0ea5e9"
        else:
            status_d0 = "BAIXA_DEMANDA"
            badge_d0 = "⚪ BAIXA DEMANDA (<20%)"
            cor_hex = "#94a3b8"

        # Saturação D-1 (Realizada ontem)
        ocupacao_d1 = round((tot_vol_d1 / capacidade_vuc) * 100, 1) if capacidade_vuc > 0 else 0
        if ocupacao_d1 > 100.0:
            status_d1 = "CRÍTICO"
            badge_d1 = "🔴 CRÍTICO (>100%)"
        elif ocupacao_d1 >= 80.0:
            status_d1 = "ALERTA"
            badge_d1 = "🟠 ALERTA (80-100%)"
        elif ocupacao_d1 >= 50.0:
            status_d1 = "EQUILIBRADO"
            badge_d1 = "🟢 EQUILIBRADO (50-80%)"
        elif ocupacao_d1 >= 20.0:
            status_d1 = "MODERADO"
            badge_d1 = "🔵 MODERADO (20-50%)"
        else:
            status_d1 = "BAIXA_DEMANDA"
            badge_d1 = "⚪ BAIXA DEMANDA (<20%)"

        # Variação D-0 vs D-1
        variacao_vol = tot_vol_d0 - tot_vol_d1
        variacao_pct = round(((tot_vol_d0 - tot_vol_d1) / tot_vol_d1 * 100), 1) if tot_vol_d1 > 0 else (100.0 if tot_vol_d0 > 0 else 0.0)

        # Monta paradas ordenadas
        sorted_stops = group.sort_values(by='seq')
        stops_list = []
        coords_path = []

        for _, stop in sorted_stops.iterrows():
            lat = float(stop['lat']) if pd.notna(stop['lat']) else None
            lng = float(stop['lng']) if pd.notna(stop['lng']) else None
            
            if lat and lng:
                coords_path.append([lat, lng])

            stops_list.append({
                "seq": int(stop['seq']),
                "iata": str(stop['md_clean']),
                "endereco": str(stop['endereco_final']),
                "bairro": str(stop.get('bairro', '') or ''),
                "cidade": str(stop.get('cidade', '') or 'São Paulo'),
                "cep": str(stop.get('cep', '') or ''),
                "janela1": str(stop.get('janela1', '') or ''),
                "janela2": str(stop.get('janela2', '') or ''),
                "horario_limite": str(stop.get('horario_limite', '') or ''),
                "vol_d0": int(stop['vol_d0']),
                "bipada_d0": int(stop['bipada_d0']),
                "backlog_d0": int(stop['backlog_d0']),
                "vol_d1": int(stop['vol_d1']),
                "sacas": int(stop['sacas']),
                "lat": lat,
                "lng": lng
            })

        rotas_list.append({
            "rota": rota_name,
            "motorista": motorista,
            "placa": placa,
            "tipo_carro": tipo_carro,
            "regiao": regiao,
            "cidade": cidade,
            "total_volume_d0": tot_vol_d0,
            "bipada_d0": tot_bip_d0,
            "backlog_d0": tot_bk_d0,
            "total_volume_d1": tot_vol_d1,
            "total_sacas": tot_sacas,
            "num_mds": num_mds,
            "ocupacao_d0_pct": ocupacao_d0,
            "ocupacao_d1_pct": ocupacao_d1,
            "status_d0": status_d0,
            "badge_d0": badge_d0,
            "status_d1": status_d1,
            "badge_d1": badge_d1,
            "cor_hex": cor_hex,
            "variacao_vol": variacao_vol,
            "variacao_pct": variacao_pct,
            "stops": stops_list,
            "coords_path": coords_path
        })

    # Ordena por maior volume planejado
    rotas_list.sort(key=lambda x: (x['total_volume_d0'], x['ocupacao_d0_pct']), reverse=True)

    # 6. Sumário Global da Frota
    total_rotas = len(rotas_list)
    rotas_criticas_d0 = len([r for r in rotas_list if r['status_d0'] == 'CRÍTICO'])
    rotas_alerta_d0 = len([r for r in rotas_list if r['status_d0'] == 'ALERTA'])
    rotas_equil_d0 = len([r for r in rotas_list if r['status_d0'] == 'EQUILIBRADO'])
    rotas_baixa_d0 = len([r for r in rotas_list if r['status_d0'] in ['MODERADO', 'BAIXA_DEMANDA']])

    rotas_criticas_d1 = len([r for r in rotas_list if r['status_d1'] == 'CRÍTICO'])
    rotas_alerta_d1 = len([r for r in rotas_list if r['status_d1'] == 'ALERTA'])

    vol_total_d0 = sum(r['total_volume_d0'] for r in rotas_list)
    vol_total_d1 = sum(r['total_volume_d1'] for r in rotas_list)

    summary = {
        "total_rotas": total_rotas,
        "total_mds_escalados": len(merged),
        "capacidade_vuc_padrao": capacidade_vuc,
        "dia_planejado": dia_planejado,
        "dia_anterior": dia_anterior,
        "volume_total_d0": vol_total_d0,
        "volume_total_d1": vol_total_d1,
        "rotas_criticas_d0": rotas_criticas_d0,
        "rotas_alerta_d0": rotas_alerta_d0,
        "rotas_equilibradas_d0": rotas_equil_d0,
        "rotas_baixa_demanda_d0": rotas_baixa_d0,
        "rotas_criticas_d1": rotas_criticas_d1,
        "rotas_alerta_d1": rotas_alerta_d1,
        "taxa_ocupacao_media_d0": round((vol_total_d0 / (total_rotas * capacidade_vuc) * 100), 1) if total_rotas > 0 else 0,
        "taxa_ocupacao_media_d1": round((vol_total_d1 / (total_rotas * capacidade_vuc) * 100), 1) if total_rotas > 0 else 0
    }

    return {
        "summary": summary,
        "rotas": rotas_list
    }


if __name__ == '__main__':
    res = calcular_saturacao_rotas()
    print("Service test successful!")
    print("Summary:", res['summary'])
    print("First route:", res['rotas'][0]['rota'], "Stops:", len(res['rotas'][0]['stops']), "Coords:", len(res['rotas'][0]['coords_path']))
